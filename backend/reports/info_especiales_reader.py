"""Lector del libro "Informe FyR Clientes Especiales <Municipio> <Comercializadora>.xlsx".

Extrae la informacion de las tres hojas de referencia:
  - ANALISIS GENERAL      : resumen mensual consumo / facturacion / recaudo
  - TARIFAS 2026          : tarifas del Acuerdo Municipal 031 de 2022
  - DESCRIPCION ESPECIALES: descripcion de los grupos de actividades economicas
"""
import unicodedata

import openpyxl

MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
         "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]


# Normaliza un texto: lo pasa a mayúsculas sin tildes para poder comparar
def _norm(v):
    if v is None:
        return ""
    s = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode()
    return s.upper().strip()


# Convierte un valor a número (quita $, comas y puntos)
def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("$", "").replace(",", "").replace(".", "").strip()
    if not s or s in ("-", "N/A", ""):
        return None
    try:
        return float(s)
    except ValueError:
        return None


# Busca una hoja por nombre (puede tener varios nombres posibles)
def _find_sheet(wb, *candidates):
    norm_candidates = tuple(_norm(c) for c in candidates)
    for name in wb.sheetnames:
        if _norm(name) in norm_candidates:
            return wb[name]
    return None


# Lee la hoja de análisis general: saca las filas mes a mes y el total/promedio
def _parse_analisis_general(ws):
    filas = []
    total = promedio = None
    for row in ws.iter_rows():
        vals = [c.value for c in row]
        a = vals[0] if vals else None
        if isinstance(a, (int, float)) and a and float(a) > 2010:
            anio, mes = int(a), 1
        else:
            anio = mes = None
            if isinstance(a, (int, float, str)):
                s = str(a).strip()
                if s.isdigit() and 2010 < int(s) <= 2100:
                    anio, mes = int(s), 1
            elif hasattr(a, "year"):
                anio = int(a.year)
                mes = int(a.month)
        if (anio and mes and 1 <= mes <= 12) or a is not None:
            label = _norm(a)
            if label in ("TOTAL", "PROMEDIO"):
                dato = {"consumo": _num(vals[1]) if len(vals) > 1 else None,
                        "facturacion": _num(vals[2]) if len(vals) > 2 else None,
                        "recaudo": _num(vals[3]) if len(vals) > 3 else None}
                if label == "TOTAL":
                    total = dato
                else:
                    promedio = dato
                continue
            if anio and mes:
                filas.append({
                    "anio": anio, "mes": mes, "mes_nombre": MESES[mes] if mes else "",
                    "consumo": _num(vals[1]) if len(vals) > 1 else None,
                    "facturacion": _num(vals[2]) if len(vals) > 2 else None,
                    "recaudo": _num(vals[3]) if len(vals) > 3 else None,
                })
    return {"filas": filas, "total": total, "promedio": promedio}


# Lee la hoja de tarifas: residenciales, no residenciales, autogeneradores,
# actividades y el "techo". Maneja dos acomodos distintos según el municipio.
def _parse_tarifas(ws):
    rows = []
    for row in ws.iter_rows():
        rows.append([c.value for c in row][:16])

    # Localizar la fila de encabezados (contiene "CLASIFICACIÓN")
    header_idx = None
    header_cols = {}
    for idx, r in enumerate(rows):
        found = {}
        for col, v in enumerate(r):
            n = _norm(v)
            if n:
                found[col] = n
        if any("CLASIFICACION" in n for n in found.values()):
            header_idx, header_cols = idx, found
            break

    residenciales, no_residenciales, autogen = [], [], []
    actividades, techo = [], None
    if header_idx is None or not header_cols:
        return {
            "residenciales": residenciales,
            "no_residenciales": no_residenciales,
            "autogeneradores": autogen,
            "actividades": actividades,
            "techo": techo,
        }

    # Layout A (Calima): tarifas mínimas urbano/rural en pesos.
    # Layout B (Andalucía): tarifas por UVT / valor en pesos / bimensual.
    layout_urbano = any("URBANO" in n for n in header_cols.values())

    data_rows = rows[header_idx + 1:]

    if layout_urbano:
        for r in data_rows:
            r = (r + [None] * 16)[:12]
            a = _norm(r[0])
            g = _norm(r[6]) if len(r) > 6 else ""

            if a.startswith("ESTRATO") or a.startswith("PARCELACION"):
                residenciales.append({
                    "clasificacion": str(r[0]).strip(),
                    "urbano": _num(r[1]),
                    "rural": _num(r[2]),
                    "bimensual_urbano": _num(r[3]) if len(r) > 3 else None,
                    "bimensual_rural": _num(r[4]) if len(r) > 4 else None,
                })

            if g.startswith("RANGO ") and len(r) > 10:
                no_residenciales.append({
                    "rango": str(r[6]).strip(),
                    "desde": str(r[7]).strip() if r[7] is not None else "",
                    "hasta": str(r[8]).strip() if r[8] is not None else "",
                    "tarifa": _num(r[9]),
                    "bimensual": _num(r[10]),
                })

            if a in ("R1", "R2", "R3", "R4") and len(r) > 4:
                h = _norm(r[2]) if r[2] is not None else ""
                if "MW" in h:
                    autogen.append({
                        "rango": a,
                        "desde": str(r[1]).strip() if r[1] is not None else "",
                        "hasta": str(r[2]).strip() if r[2] is not None else "",
                        "smlmv": _num(r[3]),
                        "tarifa": _num(r[4]),
                    })

            if g.startswith("GRUPO ") and len(r) > 9:
                actividades.append({
                    "grupo": str(r[6]).strip(),
                    "smmvl": _num(r[7]),
                    "tarifa": _num(r[8]),
                    "bimensual": _num(r[9]),
                })

            if a == "TECHO" and len(r) >= 5:
                techo = {
                    "tarifa_factor": _num(r[1]),
                    "smlmv": _num(r[2]),
                    "tarifa": _num(r[3]),
                    "bimensual": _num(r[4]),
                }
    else:
        for r in data_rows:
            r = (r + [None] * 16)[:16]
            a = _norm(r[0])
            g = _norm(r[5]) if len(r) > 5 else ""
            ac = _norm(r[11]) if len(r) > 11 else ""

            if a.startswith("ESTRATO"):
                residenciales.append({
                    "clasificacion": str(r[0]).strip(),
                    "urbano": _num(r[2]) if len(r) > 2 else None,
                    "rural": None,
                    "bimensual_urbano": _num(r[3]) if len(r) > 3 else None,
                    "bimensual_rural": None,
                })

            if g.startswith("RANGO ") and len(r) > 10:
                no_residenciales.append({
                    "rango": str(r[5]).strip(),
                    "desde": str(r[6]).strip() if r[6] is not None else "",
                    "hasta": str(r[7]).strip() if r[7] is not None else "",
                    "tarifa": _num(r[9]),
                    "bimensual": _num(r[10]),
                })

            if g == "TECHO" and len(r) >= 10:
                techo = {
                    "tarifa_factor": _num(r[6]),
                    "smlmv": _num(r[7]),
                    "tarifa": _num(r[8]),
                    "bimensual": _num(r[9]),
                }

            if ac.startswith("GRUPO ") and len(r) > 14:
                actividades.append({
                    "grupo": str(r[11]).strip(),
                    "smmvl": _num(r[12]),
                    "tarifa": _num(r[13]),
                    "bimensual": _num(r[14]),
                })

    return {
        "residenciales": residenciales,
        "no_residenciales": no_residenciales,
        "autogeneradores": autogen,
        "actividades": actividades,
        "techo": techo,
    }


_ROMAN = {n: _norm(n) for n in
          ["GRUPO I", "GRUPO II", "GRUPO III", "GRUPO IV", "GRUPO V", "GRUPO VI", "GRUPO VII"]}


# Lee la hoja de descripción: el SMMLV y los grupos de actividades económicas
def _parse_descripcion(ws):
    smmvl = None
    grupos = {}
    header_row = None
    rows = []
    for row in ws.iter_rows():
        rows.append([c.value for c in row])

    for r in rows:
        a = _norm(r[0]) if r else ""
        if a == "SMMLV" and len(r) > 1:
            smmvl = _num(r[1])
            break
    if smmvl is None:
        for r in rows:
            if len(r) > 1 and _num(r[1]) and _num(r[1]) > 1000000 and _norm(r[0]) in ("", "SMMLV"):
                cand = _num(r[1])
                if cand:
                    smmvl = cand

    for idx, r in enumerate(rows):
        labels = {}
        for col, v in enumerate(r[:8]):
            if _norm(v) in _ROMAN.values():
                labels[col] = str(v).strip()
        if len(labels) >= 2:
            header_row = idx
            break

    if header_row is not None and labels:
        for col, g in labels.items():
            items = []
            for r in rows[header_row + 1:]:
                if col < len(r) and r[col] is not None:
                    txt = str(r[col]).strip()
                    if txt and txt.upper() not in ("", "TARIFA PORCENTUAL 14% C.E.E"):
                        items.append(txt)
            grupos[g] = items

    return {"smmvl": smmvl, "grupos": grupos}


MESES_MENSUAL = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
                 "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


# Extrae la tarifa mínima, grupo y periodo de cada cuenta desde las hojas
# mensuales del libro (usa la información del mes más reciente)
def _tarifas_min_mensuales(wb):
    """Extrae la tarifa minima, grupo y periodo por ID_CUENTA desde las hojas mensuales.

    Retorna {id_cuenta_str: {"tarifa_min": float, "grupo": str, "periodo": str}}.
    Usa la información del mes más reciente disponible como referencia.
    """
    mapa = {}
    for mes in MESES_MENSUAL:
        ws = _find_sheet(wb, mes)
        if ws is None:
            continue
        for r in range(7, ws.max_row + 1):
            idv = ws.cell(r, 2).value
            if not isinstance(idv, (int, float, str)) or str(idv).strip() in ("", "None"):
                continue
            tarifa = ws.cell(r, 5).value
            grupo = ws.cell(r, 4).value
            periodo = ws.cell(r, 15).value
            if isinstance(tarifa, (int, float)):
                clave = str(int(idv)) if isinstance(idv, float) else str(idv)
                mapa[clave.strip()] = {
                    "tarifa_min": float(tarifa),
                    "grupo": str(grupo or "").strip(),
                    "periodo": str(periodo or "Mensual").strip(),
                }
    return mapa


# Parsea la hoja de cartera a 5 años (solo disponible para San Pedro).
# Cada bloque se lee en las filas fijas del libro original.
def _parse_cartera(ws):
    """Parsea la hoja CARTERA 5 AÑOS (solo San Pedro)."""
    ANIOS = [2019, 2020, 2021, 2022, 2023, 2024, 2025, 2026]
    filas = []
    total_row = None
    pagar_row = None
    variacion = {}
    retirados_2023 = []
    sin_viabilidad = []
    sin_viab_totales = None
    retirados_2025 = []

    # ---- Tabla principal: filas 5..59 --------------------------------
    # Columnas: D..K (4..11) saldos 2019..2026, L(12) VALOR TOTAL CARTERA,
    #           M(13) Notificación, N(14) SALDO
    for r in range(5, 60):
        a = ws.cell(r, 1).value
        if not (isinstance(a, (int, float)) and 1 <= a <= 55):
            break
        saldos = {}
        for i, anio in enumerate(ANIOS):
            saldos[anio] = _num(ws.cell(r, 4 + i).value)
        meses = {}
        for m in range(1, 13):
            val = _num(ws.cell(r, 14 + m).value)
            if val is not None:
                meses[str(m)] = val
        notif_raw = ws.cell(r, 13).value
        saldo_raw = ws.cell(r, 14).value
        filas.append({
            "item": int(a),
            "id": str(ws.cell(r, 2).value or "").strip(),
            "nombre": str(ws.cell(r, 3).value or "").strip(),
            "saldos": saldos,
            "meses": meses,
            "total_cartera": _num(ws.cell(r, 12).value),
            "notif": _num(notif_raw) if isinstance(notif_raw, (int, float)) else str(notif_raw).strip() if notif_raw else "",
            "saldo": _num(saldo_raw) if isinstance(saldo_raw, (int, float)) else str(saldo_raw).strip() if saldo_raw else "",
        })

    # ---- Totales principales: filas 60..61 ---------------------------
    for r in (60, 61):
        na = _norm(ws.cell(r, 1).value)
        d = {}
        for i, anio in enumerate(ANIOS):
            d[anio] = _num(ws.cell(r, 4 + i).value)
        d["total"] = _num(ws.cell(r, 12).value)
        if "A FAVOR" in na:
            total_row = d
        elif "POR PAGAR" in na:
            pagar_row = d

    # ---- Variación interanual: fila 62 (col E=2020..col K=2026) ----
    for col, anio in enumerate([2020, 2021, 2022, 2023, 2024, 2025, 2026]):
        v = _num(ws.cell(62, 5 + col).value)
        if v is not None:
            variacion[anio] = v

    # ---- Contribuyentes retirados 2023: filas 69..76 ----------------
    for r in range(69, 77):
        a = ws.cell(r, 1).value
        if not (isinstance(a, (int, float)) and 1 <= a <= 8):
            break
        retirados_2023.append({
            "item": int(a),
            "id": str(ws.cell(r, 2).value or "").strip(),
            "nombre": str(ws.cell(r, 3).value or "").strip(),
        })

    # ---- Sin viabilidad: filas 81..90 -------------------------------
    for r in range(81, 91):
        a = ws.cell(r, 1).value
        if not (isinstance(a, (int, float)) and 1 <= a <= 10):
            break
        saldos = {}
        for i, anio in enumerate([2019, 2020, 2021, 2022, 2023, 2024]):
            saldos[anio] = _num(ws.cell(r, 4 + i).value)
        sin_viabilidad.append({
            "item": int(a),
            "id": str(ws.cell(r, 2).value or "").strip(),
            "nombre": str(ws.cell(r, 3).value or "").strip(),
            "saldos": saldos,
            "total_cartera": _num(ws.cell(r, 10).value),
            "observacion": str(ws.cell(r, 12).value or "").strip(),
        })

    # ---- Totales sin viabilidad: filas 91..92 -----------------------
    sv_favor = {}
    sv_pagar = {}
    for r in (91, 92):
        d = {}
        for i, anio in enumerate([2019, 2020, 2021, 2022, 2023, 2024]):
            d[anio] = _num(ws.cell(r, 4 + i).value)
        d["total"] = _num(ws.cell(r, 10).value)
        if r == 91:
            sv_favor = d
        else:
            sv_pagar = d
    sin_viab_totales = {"favor": sv_favor, "pagar": sv_pagar}

    # ---- Contribuyentes retirados 2025: filas 96..100 ---------------
    for r in range(96, 101):
        a = ws.cell(r, 1).value
        if not (isinstance(a, (int, float)) and 1 <= a <= 5):
            break
        retirados_2025.append({
            "item": int(a),
            "id": str(ws.cell(r, 2).value or "").strip(),
            "nombre": str(ws.cell(r, 3).value or "").strip(),
            "observacion": str(ws.cell(r, 4).value or "").strip(),
        })

    return {
        "clientes": filas,
        "total_favor": total_row,
        "total_pagar": pagar_row,
        "variacion": variacion,
        "retirados_2023": retirados_2023,
        "sin_viabilidad": sin_viabilidad,
        "sin_viab_totales": sin_viab_totales,
        "retirados_2025": retirados_2025,
    }


# Función principal: abre el libro, lee cada hoja y arma el diccionario
# completo de datos que usan los informes de clientes especiales
def parse_info_especiales(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    def sheet(*names):
        for name in names:
            s = _find_sheet(wb, name)
            if s is not None:
                return s
        return None

    analisis_ws = sheet("ANALISIS GENERAL")
    tarifas_ws = sheet("TARIFAS 2026")
    desc_ws = sheet("DESCRIPCION ESPECIALES")
    cartera_ws = _find_sheet(wb, "CARTERA 5 AÑOS") or _find_sheet(wb, "CARTERA")

    acuerdo = "ACUERDO MUNICIPAL No. 031 de 2022"
    if tarifas_ws is not None:
        _v = tarifas_ws.cell(row=1, column=1).value
        if _v and "ACUERDO" in _norm(_v):
            acuerdo = str(_v).strip()

    resultado = {
        "fuente": str(path).rsplit("\\", 1)[-1].rsplit("/", 1)[-1],
        "acuerdo": acuerdo,
        "analisis_general": _parse_analisis_general(analisis_ws) if analisis_ws else {},
        "tarifas": _parse_tarifas(tarifas_ws) if tarifas_ws else {},
        "descripcion": _parse_descripcion(desc_ws) if desc_ws else {},
        "cartera": _parse_cartera(cartera_ws) if cartera_ws else None,
    }
    smmvl = resultado["descripcion"].get("smmvl")
    resultado["smmvl"] = smmvl or 1750905

    # Enriquecer clientes de la cartera con la tarifa minima mensual por cuenta
    if resultado["cartera"]:
        tar_min = _tarifas_min_mensuales(wb)
        for cl in resultado["cartera"].get("clientes", []):
            info = tar_min.get(str(cl["id"]).strip())
            cl["tarifa_min"] = (info or {}).get("tarifa_min") or 0
            cl["grupo"] = (info or {}).get("grupo") or ""
            cl["periodo"] = (info or {}).get("periodo") or "Mensual"
    return resultado


if __name__ == "__main__":
    import json
    import sys
    p = sys.argv[1]
    data = parse_info_especiales(p)
    print(json.dumps(data, ensure_ascii=False, indent=1, default=str))