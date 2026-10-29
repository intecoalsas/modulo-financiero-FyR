
import os
import json
import copy
import unicodedata
from datetime import datetime

import pandas as pd

from models import db, Municipio, Comercializadora, Report, ahora_utc
from core.utils import fmt_currency, json_safe, detect_clase, detect_sector, MESES, leer_excel
from core.depuracion import depurar_facturacion, depurar_recaudo


# Carpeta donde se guardan los Excel generados
EXPORT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "exports")


# Arma todos los datos del informe general usando los archivos de facturación y recaudo
# Con sin_depurar=True (EMSA: no hay duplicados) se toman las filas tal cual
def build_general(fact_path, rec_path, mapping_fact=None, mapping_rec=None, sin_depurar=False):
    # Primero se depuran (limpian) los dos archivos, salvo que vengan sin duplicados
    f = depurar_facturacion(fact_path, mapping_fact, sin_depurar=sin_depurar)
    r = depurar_recaudo(rec_path, mapping_rec, sin_depurar=sin_depurar)

    # Acumuladores por clase de cliente (residencial, comercial, etc.)
    # y por sector (urbano/rural) para las tablas del informe
    imp_cl = {"Residencial": 0.0, "Comercial": 0.0, "Industrial": 0.0, "Oficial": 0.0}
    cons_cl = {"Residencial": 0.0, "Comercial": 0.0, "Industrial": 0.0, "Oficial": 0.0}
    kwh_cl = {"Residencial": 0.0, "Comercial": 0.0, "Industrial": 0.0, "Oficial": 0.0}
    cnt = {"Residencial": 0, "Comercial": 0, "Industrial": 0, "Oficial": 0}
    sector_cnt = {"Urbanas": 0, "Rurales": 0}
    sector_imp = {"Urbanas": 0.0, "Rurales": 0.0}
    # Se recorre cada cuenta ya depurada y se va sumando en su casilla
    for rec in f["registros"]:
        c = rec["Clase"]
        imp_cl[c] += rec["ImpuestoAP"]
        cons_cl[c] += rec["Consumo$"]
        kwh_cl[c] += rec["ConsumokWh"]
        cnt[c] += 1
        sec = rec["Sector"]
        sector_cnt[sec] = sector_cnt.get(sec, 0) + 1
        sector_imp[sec] = sector_imp.get(sec, 0.0) + rec["ImpuestoAP"]

    # El recaudo también se acumula por clase
    rec_cl = {"Residencial": 0.0, "Comercial": 0.0, "Industrial": 0.0, "Oficial": 0.0}
    for rec in r["registros"]:
        rec_cl[rec["Clase"]] += rec["ValorAP"]

    # Consolidado: para cada clase se saca facturación vs recaudo, su diferencia
    # y qué tan eficaz fue el recaudo (qué % de lo facturado se cobró)
    consolidado = []
    totF = totR = 0.0
    for c in ["Comercial", "Industrial", "Oficial", "Residencial"]:
        fv = imp_cl[c]
        rv = rec_cl[c]
        dif = rv - fv
        efic = (rv / fv) if fv else 0
        consolidado.append([c, fv, rv, dif, efic])
        totF += fv
        totR += rv
    consolidado.append(["TOTAL", totF, totR, totR - totF, (totR / totF) if totF else 0])

    # Se cuentan cuántos registros quedaron incluidos y cuántos se descartaron
    included = f["flags"]["UNICA FACTURA"] + f["flags"]["FACTURA DIFERENTE - SUMAR TODO"]
    excluded = f["flags"]["REFACTURACION - MISMA FACTURA"] + f["flags"]["CONSUMO DUPLICADO - EXCLUIR CONSUMO"]

    # Se vuelven a leer los archivos originales (sin depurar) para comparar con lo depurado
    # (también pueden llegar directamente como DataFrames)
    orig_df = fact_path if isinstance(fact_path, pd.DataFrame) else leer_excel(fact_path)
    if mapping_fact:
        orig_df = orig_df.rename(columns=mapping_fact)
    orig_df_r = rec_path if isinstance(rec_path, pd.DataFrame) else leer_excel(rec_path, sheet_name=0)
    if mapping_rec:
        orig_df_r = orig_df_r.rename(columns=mapping_rec)

    # Tabla "clase_orig": datos del archivo original agrupados por clase
    # (registros, cuentas únicas, impuesto, consumo y kWh) más su total
    clase_orig = []
    if "desc_clase" in orig_df.columns and "Id_cliente" in orig_df.columns:
        od = orig_df.dropna(subset=["Id_cliente"]).copy()
        od["Clase"] = od["desc_clase"].map(detect_clase)
        _imp = od["monto_AP"] if "monto_AP" in od else None
        _cons = od["monto_kwh"].fillna(0) if "monto_kwh" in od else None
        _kwh = od["consumo_kwh"].fillna(0) if "consumo_kwh" in od else None
        for c in ["Residencial", "Comercial", "Industrial", "Oficial"]:
            sub = od[od["Clase"] == c]
            clase_orig.append((
                c,
                int(len(sub)),
                int(sub["Id_cliente"].nunique()),
                float(_imp[sub.index].sum()) if _imp is not None else 0.0,
                float(_cons[sub.index].sum()) if _cons is not None else 0.0,
                float(_kwh[sub.index].sum()) if _kwh is not None else 0.0,
            ))
        clase_orig.append((
            "TOTAL",
            int(len(od)),
            int(od["Id_cliente"].nunique()),
            float(_imp.sum()) if _imp is not None else 0.0,
            float(_cons.sum()) if _cons is not None else 0.0,
            float(_kwh.sum()) if _kwh is not None else 0.0,
        ))

    # Se devuelve todo listo para que la página y el Word lo pinten
    return {
        "clase_orig": clase_orig,
        "estad_general": {
            "Total registros base original": int(len(orig_df.dropna(subset=["Id_cliente"]))),
            "Usuarios \u00fanicos (ID de cuenta)": int(orig_df["Id_cliente"].nunique()),
            "Usuarios en base depurada": f["cuentas"],
            "Registros con imp. alumbrado sumado": int(len(orig_df.dropna(subset=["Id_cliente"]))),
            "Registros con consumo incluido": included,
            "Registros con consumo excluido": excluded,
        },
        "resumen_financiero": [
            ("Valor Impuesto Alumbrado P\u00fablico ($)", fmt_currency(f["orig_imp"]), fmt_currency(f["total_imp"]), fmt_currency(f["total_imp"] - f["orig_imp"])),
            ("Valor Consumo Energ\u00eda ($)", fmt_currency(f["orig_cons"]), fmt_currency(f["total_consumo"]), fmt_currency(f["total_consumo"] - f["orig_cons"])),
            ("Consumo Energ\u00eda (kWh)", "{:,.0f}".format(f["orig_kwh"]).replace(",", "."), "{:,.0f}".format(f["total_kwh"]).replace(",", "."), "{:,.0f}".format(f["total_kwh"] - f["orig_kwh"]).replace(",", ".")),
        ],
        # Explica cada "bandera": lo que se sumó y lo que no durante la depuración
        "banderas": [
            ("UNICA FACTURA", f["flags"]["UNICA FACTURA"], "SI", "SI", "Registro \u00fanico por cuenta y factura"),
            ("FACTURA DIFERENTE - SUMAR TODO", f["flags"]["FACTURA DIFERENTE - SUMAR TODO"], "SI", "SI", "Misma cuenta, facturas distintas, consumos diferentes (Regla 2)"),
            ("REFACTURACION - MISMA FACTURA", f["flags"]["REFACTURACION - MISMA FACTURA"], "SI", "NO", "Mismo ID factura: imp. SIEMPRE suma, consumo = primer valor real (Regla 1)"),
            ("CONSUMO DUPLICADO - PRIMERA OCURR.", f["flags"]["CONSUMO DUPLICADO - PRIMERA OCURR."], "SI", "SI", "Factura distinta, consumo id\u00e9ntico \u2013 primera vez suman (Regla 3)"),
            ("CONSUMO DUPLICADO - EXCLUIR CONSUMO", f["flags"]["CONSUMO DUPLICADO - EXCLUIR CONSUMO"], "SI", "NO", "Consumo id\u00e9ntico ya contabilizado: imp. suma, consumo NO (Regla 3)"),
        ],
        "clase_tabla": [
            ("Residencial", cnt["Residencial"], imp_cl["Residencial"], cons_cl["Residencial"], kwh_cl["Residencial"]),
            ("Comercial", cnt["Comercial"], imp_cl["Comercial"], cons_cl["Comercial"], kwh_cl["Comercial"]),
            ("Industrial", cnt["Industrial"], imp_cl["Industrial"], cons_cl["Industrial"], kwh_cl["Industrial"]),
            ("Oficial", cnt["Oficial"], imp_cl["Oficial"], cons_cl["Oficial"], kwh_cl["Oficial"]),
            ("TOTAL", sum(cnt.values()), sum(imp_cl.values()), sum(cons_cl.values()), sum(kwh_cl.values())),
        ],
        "sector_tabla": [
            ("Urbanas", sector_cnt["Urbanas"], sector_imp["Urbanas"]),
            ("Rurales", sector_cnt["Rurales"], sector_imp["Rurales"]),
            ("TOTAL", sum(sector_cnt.values()), sum(sector_imp.values())),
        ],
        "estad_recaudo": {
            "Total registros base original": int(len(orig_df_r.dropna(subset=["Id_cliente"]))),
            "Usuarios \u00fanicos (ID de cuenta)": int(orig_df_r["Id_cliente"].nunique()),
            "Usuarios en base depurada": r["cuentas"],
            "Registros \u00fanicos (sin duplicar)": r["flags"]["UNICO"],
            "Registros BASE (con duplicados)": r["flags"]["BASE"],
            "Registros DUPLICADOS acumulados": r["flags"]["DUPLICADO"],
        },
        "resumen_recaudo": [
            ("Valor Recaudo A.P ($)", fmt_currency(r["orig_total"]), fmt_currency(r["total_valor"]), fmt_currency(r["total_valor"] - r["orig_total"])),
        ],
        "consolidado": consolidado,
        "facturacion": f["registros"],
        "recaudo": r["registros"],
        "f": f,
        "r": r,
    }


# Lista fija de clientes especiales (cuentas grandes o de servicios públicos)
# que se usan cuando no hay archivos cargados. Cada uno es:
# (ID, nombre, grupo, tarifa mínima, consumo, facturación, recaudo, periodicidad)
ESPECIALES = [
    ("6147040000", "SERVICIOS AUXILIAR SUBESTACION DARIEN", "AUTOGENERADOR R4", 35018100, 3087206, 5252715, 5252715, "Mensual"),
    ("406900000", "RIVERA, JORGUE ELIECER", "GRUPO 1", 1750905, 0, 0, 0, "Mensual"),
    ("5177950000", "WISP TELECOMUNICACIONES S.A.S", "GRUPO 1", 1750905, 0, 0, 0, "Mensual"),
    ("3840800000", "CABLE ANTENA LTDA", "GRUPO 1", 1750905, 0, 0, 0, "Mensual"),
    ("1964300000", "INVERCER SAS", "GRUPO 2", 2626358, 0, 0, 0, "Mensual"),
    ("4926800000", "RODRIGUEZ ASCUNTAR, ADRIANA NORELLA", "GRUPO 2", 0, 2739845, 5252716, 0, "Bimestral Impar"),
    ("4911400000", "OSORIO VASCO, CAMILO ANDRES", "GRUPO 2", 2626358, 2942586, 2637389, 2635813, "Mensual"),
    ("1847040000", "COMERCIAL AVANADE S.A.S.", "GRUPO 2", 2626358, 2677583, 2626358, 5254292, "Mensual"),
    ("5507150000", "GASES DE OCCIDENTE, S.A.E.S.P", "GRUPO 2", 2626358, 615197, 2626358, 5255868, "Mensual"),
    ("9498800000", "PLANTA DE ACUEDUCTO", "GRUPO 2", 0, 8687763, 60376, 60376, "Bimestral Impar"),
    ("1039940000", "DISTRACOM S A", "GRUPO 2", 2626358, 3160200, 30188, 30188, "Mensual"),
    ("1699400000", "TRANS CALIMA", "GRUPO 2", 2626358, 112426, 6943, 6943, "Mensual"),
    ("2808560000", "EMCALIMA ESP", "GRUPO 2", 2626358, 30662, 6943, 0, "Mensual"),
    ("9803900000", "TRANSPORTES LOPEZ S A, ALFONSO", "GRUPO 2", 2626358, 0, 0, 0, "Mensual"),
    ("1143300000", "ATC SITIOS DE COLOMBIA S.A.S.", "GRUPO 3", 0, 1740361, 10505430, 0, "Bimestral Impar"),
    ("5510400000", "BANCO AGRARIO DE COLOMBIA", "GRUPO 3", 5252715, 361554, 5252715, 5252715, "Mensual"),
    ("7721820000", "BANCO POPULAR", "GRUPO 3", 5252715, 126739, 6943, 19055, "Mensual"),
    ("2407800000", "ALIANZA FIDUCIARIA S A", "GRUPO 3", 10505430, 0, 0, 0, "Bimestral Par"),
    ("821620000", "COLOMBIA TELECOMUNICACIONES, S.A", "GRUPO 4", 7879073, 2723867, 5252715, 10527491, "Mensual"),
    ("2788340000", "COLOMBIA MOVIL S.A.", "GRUPO 4", 7879073, 0, 0, 0, "Mensual"),
    ("1264150000", "COLOMBIA TELECOMUNICACIONES, S.A /RADIO NET SOLUCIONES S.A", "GRUPO 4", 0, 2325303, 15758146, 15758146, "Bimestral Impar"),
    ("5819750000", "COLOMBIA TELECOMUNICACIONES, S.A / RADIO NET SOLUCIONES S.A", "GRUPO 4", 0, 287362, 15758146, 0, "Bimestral Impar"),
    ("3744040000", "COMCEL S.A.", "GRUPO 4", 0, 751406, 15758146, 15758146, "Bimestral Impar"),
    ("2074300000", "COLOMBIA TELECOMUNICACIONES, S.A. ESP.", "GRUPO 4", 7879073, 2908189, 5252715, 0, "Mensual"),
    ("6860650000", "CVC EDIFICIO PRINCIPAL DE LA", "GRUPO 4", 7879073, 249010, 7879073, 7879073, "Mensual"),
    ("2905900000", "CVC PTO VIGILANCIA", "GRUPO 4", 0, 38726, 13886, 13886, "Bimestral Impar"),
    ("8950850000", "RADIO NET SOLUCIONES S.A", "GRUPO 4", 7879073, 645148, 15051, 30139, "Mensual"),
    ("8236396730", "PARTNERS TELECOM COLOMBIA", "GRUPO 4", 7879073, 1020196, 30188, 0, "Mensual"),
    ("7110400000", "INCIVA", "GRUPO 4", 7879073, 538436, 15044, 13607, "Mensual"),
    ("6222650000", "AZTECA COMUNICACIONES, COLOMBIA S.A.S", "GRUPO 4", 7879073, 250837, 8906, 20981, "Mensual"),
    ("5500000000", "COMFANDI", "GRUPO 5", 21010860, 42607654, 21010860, 21010860, "Mensual"),
    ("1901900000", "CARTON DE COLOMBIA S.A.", "GRUPO 5", 0, 0, 0, 0, "Bimestral Impar"),
    ("7740800000", "CALIMA RESORT", "GRUPO 5", 0, 0, 0, 0, "Bimestral Impar"),
    ("3012800000", "EPSA CASINO MADRONAL", "GRUPO 6", 0, 1076595, 29886, 29886, "Bimestral Impar"),
    ("7147040000", "SERVICIOS AUXILIARES SUBESTACION, RIO BRAVO", "GRUPO 6", 35018100, 0, 0, 0, "Mensual"),
    ("7522900000", "TESTIGOS DE JEHOVA", "GRUPO 7", 4377, 58536, 8205, 5426, "Mensual"),
    ("7700400000", "SALON CLUB DE LEONES", "GRUPO 7", 4377, 20441, 4377, 4377, "Mensual"),
    ("1563160000", "IGLESIA PENTECOSTAL UNIDA DE, COLOMBIA", "GRUPO 7", 4377, 270380, 37934, 27553, "Mensual"),
    ("5911400000", "CUERPO BOMBEROS VOLUNTARIOS", "GRUPO 7", 4377, 275955, 38715, 33885, "Mensual"),
    ("9700400000", "CASA PARROQUIA", "GRUPO 7", 4377, 36236, 5128, 5089, "Mensual"),
    ("6600400000", "DEFENSA CIVIL DAR", "GRUPO 7", 4377, 225781, 31609, 0, "Mensual"),
    ("6700400000", "CLUB DE LEONES", "GRUPO 7", 4377, 27874, 4440, 0, "Mensual"),
    ("8664240000", "PARROQUIA DEL PERP SOCORRO DEL DARIEN", "GRUPO 7", 4377, 253655, 35512, 35512, "Mensual"),
    ("8700400000", "IGLESIA PARROQUIAL", "GRUPO 7", 4377, 346569, 48520, 48520, "Mensual"),
    ("4131840000", "IGLESIA PETENCOSTAL", "GRUPO 7", 0, 0, 0, 0, "Bimestral Impar"),
    ("7086520000", "IGLESIA CRISTIANA PE", "GRUPO 7", 4377, 0, 0, 0, "Mensual"),
    ("2920400000", "CASA DEL ANCIANO", "GRUPO 7", 4377, 0, 0, 0, "Mensual"),
]


# Arma los datos del informe de clientes especiales.
# Si se pasan los archivos y la comercializadora, los calcula desde ahí;
# si no, usa la lista fija ESPECIALES.
def build_especiales(mes=None, comercializadora=None, anio=None,
                     fact_path=None, rec_path=None,
                     mapping_fact=None, mapping_rec=None):
    if comercializadora and fact_path and rec_path and anio:
        result = build_especiales_desde_archivos(
            mes, anio, comercializadora, fact_path, rec_path,
            mapping_fact=mapping_fact, mapping_rec=mapping_rec)
        if result is not None:
            return result
    # TEC = tope máximo del impuesto que se puede cobrar
    TEC = 35018100
    filas = []
    for e in ESPECIALES:
        idc, nombre, clas, tarifa, consumo, fact, rec, periodo = e
        saldo = rec - fact

        # El impuesto (IAP) se calcula así: mínimo la tarifa, pero al menos
        # el 14% del consumo, y nunca pasarse del tope TEC
        calc_iap = min(max(tarifa, consumo * 0.14), TEC)
        saldo_fl = fact - calc_iap
        pct_cee = (fact / consumo) if consumo > 0 else 0.0
        val_rec = rec - calc_iap

        eps = TEC * 0.001
        # Si lo facturado es igual al impuesto calculado, aplicó la Tarifa Especial
        if abs(fact - calc_iap) < eps or clas in ("GRUPO 5", "GRUPO 7"):
            val_tarifa_min = "Tarifa Especial"
        else:
            val_tarifa_min = "Otra Tarifa"

        # Según qué % del consumo se facturó, se clasifica el tipo de tarifa
        if abs(pct_cee - 0.14) < 0.01 and fact >= consumo * 0.14:
            val_tarifa_pct = "Especial"
        elif pct_cee < 0.10 and fact > 0:
            val_tarifa_pct = "Residencial"
        else:
            val_tarifa_pct = "Otra Tarifa"

        # Revisa si la cuenta debía facturar este mes (según si es bimestral o no)
        # y si no lo hizo, lo deja marcado para que se revise
        if fact == 0:
            estado_fact = "SIN FACTURACION"
            if "Bimestral Par" in periodo and mes and mes % 2 == 1:
                estado_fact = "NO CORRESPONDE BIMESTRAL (par - mes impar)"
            elif "Bimestral Impar" in periodo and mes and mes % 2 == 0:
                estado_fact = "NO CORRESPONDE BIMESTRAL (impar - mes par)"
            elif "Bimestral Impar" in periodo and mes and mes % 2 == 1:
                estado_fact = "SIN FACTURAR (debi\u00f3 facturar mes impar) - REVISAR"
            elif "Bimestral Par" in periodo and mes and mes % 2 == 0:
                estado_fact = "SIN FACTURAR (debi\u00f3 facturar mes par) - REVISAR"
            elif "Bimestral" in periodo:
                estado_fact = "SIN FACTURAR (bimestral) - REVISAR PERIODICIDAD"
        else:
            estado_fact = "FACTURADO"

        # Una fila por cada cliente especial con todos los datos calculados
        filas.append({
            "id": idc, "nombre": nombre, "clase": clas,
            "tarifa_minima": tarifa, "consumo": consumo,
            "facturacion": fact, "recaudo": rec,
            "verificado": "SI" if fact != 0 else "NO", "periodo": periodo,
            "calc_iap": calc_iap, "saldo_fl": saldo_fl,
            "pct_cee": pct_cee, "val_rec": val_rec,
            "val_tarifa_pct": val_tarifa_pct, "val_tarifa_min": val_tarifa_min,
            "estado_fact": estado_fact,
        })
    # Se agrupan por clase (grupo) para el consolidado
    por_clase = {}
    for e in filas:
        por_clase.setdefault(e["clase"], {"n": 0, "fact": 0.0, "rec": 0.0, "consumo": 0.0, "saldo_fl": 0.0, "saldo_pos": 0.0, "saldo_neg": 0.0})
        por_clase[e["clase"]]["n"] += 1
        por_clase[e["clase"]]["fact"] += e["facturacion"]
        por_clase[e["clase"]]["rec"] += e["recaudo"]
        por_clase[e["clase"]]["consumo"] += e["consumo"]
        por_clase[e["clase"]]["saldo_fl"] += e["saldo_fl"]
        if e["saldo_fl"] > 0:
            por_clase[e["clase"]]["saldo_pos"] += e["saldo_fl"]
        else:
            por_clase[e["clase"]]["saldo_neg"] += e["saldo_fl"]

    # Totales generales de todos los clientes especiales
    total_fact = sum(e["facturacion"] for e in filas)
    total_rec = sum(e["recaudo"] for e in filas)
    total_saldo_fl = sum(e["saldo_fl"] for e in filas)
    total_val_rec = sum(e["val_rec"] for e in filas)
    saldo_pos = sum(e["saldo_fl"] for e in filas if e["saldo_fl"] > 0)
    saldo_neg = sum(e["saldo_fl"] for e in filas if e["saldo_fl"] < 0)

    # Cuánto impuesto se dejó de facturar y cuánto no llegó al municipio
    iap_dejado_facturar = sum(max(0.0, e["calc_iap"] - e["facturacion"]) for e in filas)
    iap_no_percibido_municipio = sum(max(0.0, e["calc_iap"] - e["recaudo"]) for e in filas)

    # Listas de interés para el análisis: los que no facturaron,
    # los que usan otra tarifa y los bimestrales
    sin_facturar = [e for e in filas if e["facturacion"] == 0]
    otras_tarifas = [e for e in filas if e["val_tarifa_min"] == "Otra Tarifa"]
    bimestrales = [e for e in filas if "Bimestral" in e["periodo"]]

    analisis = {
        "mes": mes,
        "iap_dejado_facturar": iap_dejado_facturar,
        "iap_no_percibido_municipio": iap_no_percibido_municipio,
        "cuentas_sin_facturar": len(sin_facturar),
        "cuentas_otras_tarifas": len(otras_tarifas),
        "cuentas_bimestrales": len(bimestrales),
        "list_sin_facturar": _compact_analisis(sin_facturar),
        "list_otras_tarifas": _compact_analisis(otras_tarifas),
        "list_bimestrales": _compact_analisis(bimestrales),
    }

    return {
        "filas": filas,
        "por_clase": por_clase,
        "grupos": sorted(por_clase.keys(), key=lambda x: (x != "AUTOGENERADOR R4", x)),
        "totales": {
            "consumo": sum(e["consumo"] for e in filas),
            "fact": total_fact,
            "rec": total_rec,
            "tarifa": sum(e["tarifa_minima"] for e in filas),
            "saldo": total_rec - total_fact,
            "cuentas": len(filas),
            "eficacia": (total_rec / total_fact) if total_fact else 0,
            "total_saldo_fl": total_saldo_fl,
            "total_val_rec": total_val_rec,
            "saldo_pos": saldo_pos,
            "saldo_neg": saldo_neg,
        },
        "analisis": analisis,
    }


# Deja la lista "compacta" (solo datos importantes) para mostrar en el Word
def _compact_analisis(filas):
    return [{
        "id": e["id"], "nombre": e["nombre"], "clase": e["clase"],
        "consumo": e["consumo"], "facturacion": e["facturacion"],
        "calc_iap": e["calc_iap"], "periodo": e["periodo"],
        "estado_fact": e["estado_fact"],
    } for e in filas]


# Igual que build_especiales pero calculando con los archivos reales que se subieron.
# Usa la cartera guardada de la comercializadora para saber las tarifas de cada cuenta.
def build_especiales_desde_archivos(mes, anio, comercializadora, fact_path, rec_path,
                                    mapping_fact=None, mapping_rec=None):
    TEC = 35018100

    # Trae la cartera (clientes con sus tarifas) del año indicado
    cartera = get_cartera(comercializadora, anio)
    if not cartera:
        return None

    # Solo se tienen en cuenta los clientes que tienen tarifa mínima (están activos)
    clientes_activos = [
        c for c in cartera.get("clientes", [])
        if (c.get("tarifa_min") or 0) > 0
    ]
    if not clientes_activos:
        return None

    ids_ref = {str(c["id"]).strip(): c for c in clientes_activos}

    import warnings
    warnings.simplefilter("ignore")

    # Se leen los archivos y se aplica el mapeo de columnas (si el usuario
    # tuvo que decir cómo se llaman las columnas, aquí se renombran antes
    # de usarlas; sin esto revienta con KeyError 'Id_cliente')
    df_f = leer_excel(fact_path)
    if mapping_fact:
        df_f = df_f.rename(columns=mapping_fact)
    df_f["id_s"] = df_f["Id_cliente"].apply(lambda v: str(int(v)) if pd.notna(v) else "")

    df_r = leer_excel(rec_path)
    if mapping_rec:
        df_r = df_r.rename(columns=mapping_rec)
    df_r["id_s"] = df_r["Id_cliente"].apply(lambda v: str(int(v)) if pd.notna(v) else "")

    # Los archivos mandan en el nombre: se toma el NOMBRE_USUARIO del archivo
    # de facturación (y si no está, del de recaudo). La cartera solo aporta
    # grupo, tarifa y periodo, que se mantienen con edición manual.
    def _col_nombre(df):
        for c in df.columns:
            if str(c).strip().lower() == "nombre":
                return c
        return None

    nom_f = {}
    _cf = _col_nombre(df_f)
    if _cf:
        for gid, grp in df_f.groupby("id_s"):
            vals = grp[_cf].dropna().astype(str).str.strip()
            vals = vals[vals != ""]
            if len(vals):
                nom_f[gid] = vals.mode().iloc[0] if not vals.mode().empty else vals.iloc[0]
    nom_r = {}
    _cr = _col_nombre(df_r)
    if _cr:
        for gid, grp in df_r.groupby("id_s"):
            vals = grp[_cr].dropna().astype(str).str.strip()
            vals = vals[vals != ""]
            if len(vals):
                nom_r[gid] = vals.mode().iloc[0] if not vals.mode().empty else vals.iloc[0]

    # Por cada cuenta de la cartera se suma lo facturado y el consumo
    # (el consumo solo se cuenta una vez por factura)
    por_id = {}
    for gid, grp in df_f.groupby("id_s"):
        if gid not in ids_ref:
            continue
        d = por_id.setdefault(gid, {"consumo": 0.0, "fact": 0.0})
        for _, row in grp.iterrows():
            d["fact"] += float(row.get("monto_AP") or 0)
        nf = grp.drop_duplicates("numero_factura")
        d["consumo"] += nf["monto_kwh"].sum()

    # Y también el total recaudado por cuenta
    rec_map = {}
    for gid, grp in df_r.groupby("id_s"):
        if gid in ids_ref:
            rec_map[gid] = grp["valor_recaudado"].sum()

    filas = []
    for cid, ref in ids_ref.items():
        # El nombre sale del archivo del mes (los archivos mandan);
        # la cartera solo es respaldo si el ID no vino en los archivos
        nombre = nom_f.get(cid) or nom_r.get(cid) or ref.get("nombre", "")
        clas = ref.get("grupo", "")
        tarifa = float(ref.get("tarifa_min") or 0)
        periodo = ref.get("periodo") or "Mensual"
        es_bimensual = "Bimensual" in periodo

        d = por_id.get(cid, {})
        consumo = d.get("consumo", 0.0)
        fact = d.get("fact", 0.0)
        rec = rec_map.get(cid, 0.0)

        # A los bimestrales no se les calcula impuesto este mes (no les corresponde)
        if es_bimensual:
            calc_iap = 0.0
        else:
            calc_iap = min(max(tarifa, consumo * 0.10), TEC)

        saldo_fl = fact - calc_iap
        pct_cee = (fact / consumo) if consumo > 0 else 0.0
        val_rec = rec - calc_iap

        # Clasifica la tarifa según la diferencia entre facturado y calculado
        if abs(saldo_fl) <= 10000:
            val_tarifa_min = "Tarifa Especial"
        else:
            val_tarifa_min = "Otra Tarifa"

        # Según el % del consumo que se facturó se decide el tipo de tarifa
        if 0.05 <= pct_cee < 0.06:
            val_tarifa_pct = "Residencial"
        elif 0.08 <= pct_cee < 0.09:
            val_tarifa_pct = "No Residencial"
        elif 0.09 <= pct_cee < 0.11:
            val_tarifa_pct = "Especial"
        else:
            val_tarifa_pct = "Otra Tarifa"

        verificado = "SI" if fact != 0 else "NO"

        # Igual que en build_especiales: revisa si la cuenta debía facturar este mes
        if fact == 0:
            estado_fact = "SIN FACTURACION"
            if "Bimestral Par" in periodo and mes and mes % 2 == 1:
                estado_fact = "NO CORRESPONDE BIMESTRAL (par - mes impar)"
            elif "Bimestral Impar" in periodo and mes and mes % 2 == 0:
                estado_fact = "NO CORRESPONDE BIMESTRAL (impar - mes par)"
            elif "Bimestral Impar" in periodo and mes and mes % 2 == 1:
                estado_fact = "SIN FACTURAR (debi\u00f3 facturar mes impar) - REVISAR"
            elif "Bimestral Par" in periodo and mes and mes % 2 == 0:
                estado_fact = "SIN FACTURAR (debi\u00f3 facturar mes par) - REVISAR"
            elif "Bimestral" in periodo:
                estado_fact = "SIN FACTURAR (bimestral) - REVISAR PERIODICIDAD"
        else:
            estado_fact = "FACTURADO"

        filas.append({
            "id": cid, "nombre": nombre, "clase": clas,
            "tarifa_minima": tarifa, "consumo": consumo,
            "facturacion": fact, "recaudo": rec,
            "verificado": verificado, "periodo": periodo,
            "calc_iap": calc_iap, "saldo_fl": saldo_fl,
            "pct_cee": pct_cee, "val_rec": val_rec,
            "val_tarifa_pct": val_tarifa_pct, "val_tarifa_min": val_tarifa_min,
            "estado_fact": estado_fact,
        })

    # Se ordenan por grupo y nombre para que se vea ordenado
    filas.sort(key=lambda x: (x["clase"], x["nombre"]))

    # El resto es igual a build_especiales: agrupar, totalizar y devolver
    por_clase = {}
    for e in filas:
        por_clase.setdefault(e["clase"], {"n": 0, "fact": 0.0, "rec": 0.0, "consumo": 0.0, "saldo_fl": 0.0, "saldo_pos": 0.0, "saldo_neg": 0.0})
        por_clase[e["clase"]]["n"] += 1
        por_clase[e["clase"]]["fact"] += e["facturacion"]
        por_clase[e["clase"]]["rec"] += e["recaudo"]
        por_clase[e["clase"]]["consumo"] += e["consumo"]
        por_clase[e["clase"]]["saldo_fl"] += e["saldo_fl"]
        if e["saldo_fl"] > 0:
            por_clase[e["clase"]]["saldo_pos"] += e["saldo_fl"]
        else:
            por_clase[e["clase"]]["saldo_neg"] += e["saldo_fl"]

    total_fact = sum(e["facturacion"] for e in filas)
    total_rec = sum(e["recaudo"] for e in filas)
    total_saldo_fl = sum(e["saldo_fl"] for e in filas)
    total_val_rec = sum(e["val_rec"] for e in filas)
    saldo_pos = sum(e["saldo_fl"] for e in filas if e["saldo_fl"] > 0)
    saldo_neg = sum(e["saldo_fl"] for e in filas if e["saldo_fl"] < 0)

    iap_dejado_facturar = sum(max(0.0, e["calc_iap"] - e["facturacion"]) for e in filas)
    iap_no_percibido_municipio = sum(max(0.0, e["calc_iap"] - e["recaudo"]) for e in filas)

    sin_facturar = [e for e in filas if e["facturacion"] == 0]
    otras_tarifas = [e for e in filas if e["val_tarifa_min"] == "Otra Tarifa"]
    bimestrales = [e for e in filas if "Bimensual" in e["periodo"]]

    analisis = {
        "mes": mes,
        "iap_dejado_facturar": iap_dejado_facturar,
        "iap_no_percibido_municipio": iap_no_percibido_municipio,
        "cuentas_sin_facturar": len(sin_facturar),
        "cuentas_otras_tarifas": len(otras_tarifas),
        "cuentas_bimestrales": len(bimestrales),
        "list_sin_facturar": _compact_analisis(sin_facturar),
        "list_otras_tarifas": _compact_analisis(otras_tarifas),
        "list_bimestrales": _compact_analisis(bimestrales),
    }

    return {
        "filas": filas,
        "por_clase": por_clase,
        "grupos": sorted(por_clase.keys(), key=lambda x: (x != "AUTOGENERADOR R4", x)),
        "totales": {
            "consumo": sum(e["consumo"] for e in filas),
            "fact": total_fact,
            "rec": total_rec,
            "tarifa": sum(e["tarifa_minima"] for e in filas),
            "saldo": total_rec - total_fact,
            "cuentas": len(filas),
            "eficacia": (total_rec / total_fact) if total_fact else 0,
            "total_saldo_fl": total_saldo_fl,
            "total_val_rec": total_val_rec,
            "saldo_pos": saldo_pos,
            "saldo_neg": saldo_neg,
        },
        "analisis": analisis,
    }


# Le pone el formato (colores, bordes) a las hojas de los Excel exportados
def _estilo(ws):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9D9D9")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    # Primera fila (los títulos de columna) con color y negrita
    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = bdr
    # El resto de celdas solo con borde
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = bdr
    # La primera fila queda fija al hacer scroll
    ws.freeze_panes = "A2"


# Crea el Excel del informe general (varias hojas) y lo guarda en exports
def exportar_general(data, filename):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"
    ws.append(["INDICADOR", "VALOR"])
    for k, v in data["estad_general"].items():
        ws.append([k, v])
    ws.append([])
    ws.append(["CONCEPTO", "BASE ORIGINAL", "BASE DEPURADA", "DIFERENCIA"])
    for r in data["resumen_financiero"]:
        ws.append(list(r))
    ws.append([])
    ws.append(["BANDERA", "CANTIDAD", "IMP.", "CONS.", "DESCRIPCION"])
    for b in data["banderas"]:
        ws.append(list(b))
    _estilo(ws)

    ws2 = wb.create_sheet("FACTURACION")
    ws2.append(["ID", "EMPRESA", "NOMBRE", "DIRECCION", "CLASE", "TIPO", "ESTRATO", "SECTOR", "PERIODICIDAD", "IMP AP", "CONSUMO $", "KWH"])
    for rec in data["facturacion"]:
        ws2.append([rec["Id_cliente"], rec["Empresa"], rec["Nombre"], rec["Direccion"], rec["Clase"], rec["Tipo"], rec["Estrato"], rec["Sector"], rec["Periodicidad"], rec["ImpuestoAP"], rec["Consumo$"], rec["ConsumokWh"]])
    _estilo(ws2)

    ws3 = wb.create_sheet("RECAUDO")
    ws3.append(["ID", "EMPRESA", "TIPO", "NOMBRE", "DIRECCION", "CLASE", "SECTOR", "ESTRATO", "VALOR AP DEP"])
    for rec in data["recaudo"]:
        ws3.append([rec["Id_cliente"], rec["Empresa"], rec["Tipo"], rec["Nombre"], rec["Direccion"], rec["Clase"], rec["Sector"], rec["Estrato"], rec["ValorAP"]])
    _estilo(ws3)

    ws4 = wb.create_sheet("DINAMICA")
    ws4.append(["CLASIFICACION", "FACTURACION", "RECAUDO", "DIFERENCIA", "EFICACIA"])
    for r in data["consolidado"]:
        ws4.append(r)
    _estilo(ws4)

    path = os.path.join(EXPORT_DIR, filename)
    wb.save(path)
    return path


# Crea el Excel del informe de especiales y lo guarda en exports
def exportar_especiales(data, filename):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "ESPECIALES"
    ws.append(["ID_CUENTA", "NOMBRE", "CLASIFICACION", "TARIFA MINIMA", "CONSUMO", "FACTURACION", "RECAUDO", "VERIF. FACTURADO", "PERIODO",
               "CALC IAP", "SALDOS F-L", "% CEE", "VALIDACION REC", "VALIDACION %", "VALIDACION MIN"])
    for e in data["filas"]:
        ws.append([e["id"], e["nombre"], e["clase"], e["tarifa_minima"], e["consumo"], e["facturacion"], e["recaudo"], e["verificado"], e["periodo"],
                   e["calc_iap"], e["saldo_fl"], e["pct_cee"], e["val_rec"], e["val_tarifa_pct"], e["val_tarifa_min"]])
    _estilo(ws)

    ws2 = wb.create_sheet("DASHBOARD")
    ws2.append(["GRUPO", "CUENTAS", "FACTURACION", "RECAUDO", "CONSUMO"])
    for g in data["grupos"]:
        d = data["por_clase"][g]
        ws2.append([g, d["n"], d["fact"], d["rec"], d["consumo"]])
    _estilo(ws2)

    path = os.path.join(EXPORT_DIR, filename)
    wb.save(path)
    return path


# Trata de averiguar el municipio, la comercializadora, el año y el mes
# mirando los datos dentro de los archivos de facturación y recaudo
def detectar_metadatos(fact_path, rec_path, mapping_fact=None, mapping_rec=None):
    df_f = leer_excel(fact_path, sheet_name=0)
    df_r = leer_excel(rec_path, sheet_name=0)
    if mapping_fact:
        df_f = df_f.rename(columns=mapping_fact)
    if mapping_rec:
        df_r = df_r.rename(columns=mapping_rec)

    municipio = ""
    # El municipio se busca primero en facturación y si no está, en recaudo
    if "municipio" in df_f.columns:
        vals = df_f["municipio"].dropna().astype(str).str.strip()
        if len(vals):
            municipio = vals.mode().iloc[0] if not vals.mode().empty else vals.iloc[0]
    if not municipio and "municipio" in df_r.columns:
        vals = df_r["municipio"].dropna().astype(str).str.strip()
        if len(vals):
            municipio = vals.mode().iloc[0] if not vals.mode().empty else vals.iloc[0]

    comercializadora = ""
    # Se busca el nombre que contenga "CELSIA" en cualquier columna de nombre
    for col in ("Nombre", "nombre"):
        for df in (df_f, df_r):
            if col in df.columns:
                hits = df[col].dropna().astype(str)
                celsia = hits[hits.str.upper().str.contains("CELSIA", na=False)]
                if len(celsia):
                    comercializadora = celsia.iloc[0]
                    break
            if comercializadora:
                break
        if comercializadora:
            break

    # El año y el mes se sacan de la fecha que tenga el archivo
    anio = mes = None
    for col in ("fecha_generacion", "fecha_vencimiento", "fecha_aplicacion", "fecha_pago"):
        for df in (df_f, df_r):
            if col in df.columns:
                fechas = pd.to_datetime(df[col], errors="coerce").dropna()
                if len(fechas):
                    anio = int(fechas.dt.year.mode().iloc[0])
                    mes = int(fechas.dt.month.mode().iloc[0])
                    break
            if anio is not None:
                break
        if anio is not None:
            break

    municipio = municipio.strip().title()
    comercializadora = comercializadora.replace(".", "").strip().title()

    # Si el municipio o la comercializadora ya existen en la base de datos,
    # se usa el nombre que ya está guardado (para no crear duplicados)
    if municipio:
        existentes = [m.nombre for m in Municipio.query.filter(
            Municipio.nombre.ilike(f"%{municipio}%"))]
        if existentes:
            municipio = existentes[0]
    if comercializadora:
        if "CELSIA" in comercializadora.upper():
            existentes = [c.nombre for c in Comercializadora.query.filter(
                Comercializadora.nombre.ilike("Celsia"))]
            comercializadora = existentes[0] if existentes else "Celsia"
        else:
            existentes = [c.nombre for c in Comercializadora.query.filter(
                Comercializadora.nombre.ilike(f"%{comercializadora}%"))]
            if existentes:
                comercializadora = existentes[0]

    return municipio, comercializadora, anio, mes


# Nombre normalizado para comparar sin tildes, mayúsculas ni espacios de más.
# Así "Andalucia" y "Andalucía" son el mismo municipio y no se duplica.
def _norm_nombre(s):
    return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().strip().lower()


# Busca un municipio por nombre y si no existe lo crea.
# Si ya existe uno igual sin importar tildes/mayúsculas, usa ese (no duplica).
def get_or_create_municipio(nombre):
    nombre = (nombre or "").strip()
    m = Municipio.query.filter_by(nombre=nombre).first()
    if not m:
        objetivo = _norm_nombre(nombre)
        for cand in Municipio.query.all():
            if _norm_nombre(cand.nombre) == objetivo:
                m = cand
                break
    if not m:
        m = Municipio(nombre=nombre)
        db.session.add(m)
        db.session.flush()
    return m


# Busca una comercializadora en el municipio y si no existe la crea.
# Igual que municipios: no duplica por tildes/mayúsculas.
def get_or_create_comerc(municipio_id, nombre):
    nombre = (nombre or "").strip()
    c = Comercializadora.query.filter_by(municipio_id=municipio_id, nombre=nombre).first()
    if not c:
        objetivo = _norm_nombre(nombre)
        for cand in Comercializadora.query.filter_by(municipio_id=municipio_id).all():
            if _norm_nombre(cand.nombre) == objetivo:
                c = cand
                break
    if not c:
        c = Comercializadora(municipio_id=municipio_id, nombre=nombre)
        db.session.add(c)
        db.session.flush()
    return c


# Devuelve la info de clientes especiales de una comercializadora.
# Si se pasa un año, busca ese año; si no, el más reciente
def get_info_especiales(comercializadora, anio=None):
    if not comercializadora or not comercializadora.info_especiales:
        return None
    d = comercializadora.info_especiales
    if not isinstance(d, dict):
        return None
    if anio is not None and str(anio) in d:
        return d[str(anio)]
    keys = [k for k in d if str(k).isdigit()]
    if keys:
        return d[sorted(keys, key=int, reverse=True)[0]]
    return None


# Devuelve las "carpetas" (años) del análisis general de una comercializadora
def get_analisis_general_folders(comercializadora):
    if not comercializadora or not comercializadora.analisis_general:
        return {}
    d = comercializadora.analisis_general
    if not isinstance(d, dict):
        return {}
    return d


# Guarda o actualiza el consumo, facturación y recaudo de un mes/año
# en el análisis general de la comercializadora. Para Villavicencio/EMSA
# también se guarda el desglose FYR del mes (fact/rec/base/kwh por
# clasificación y estrato) para poder graficar el acumulado anual.
def upsert_analisis_general(comercializadora, anio, mes, consumo, facturacion, recaudo,
                            fyr_data=None, commit=True):
    if not comercializadora or not anio or not mes:
        return
    folders = dict(comercializadora.analisis_general or {})
    year = dict(folders.get(str(anio), {}) or {})
    mes_data = {
        "consumo": consumo or 0,
        "facturacion": facturacion or 0,
        "recaudo": recaudo or 0,
    }
    if fyr_data:
        mes_data["fyr"] = fyr_data
    year[str(mes)] = mes_data
    folders[str(anio)] = year
    comercializadora.analisis_general = folders
    if commit:
        db.session.commit()


# Suma el desglose FYR de todos los meses cargados de un año
# (solo Villavicencio/EMSA guardan fyr por mes). Devuelve
# {"Residencial – Estrato 1": {"fact": x, "rec": y, "base": z, "kwh": w}, ...}
def agregar_fyr_anual(analisis_general, anio):
    year = (analisis_general or {}).get(str(anio), {}) or {}
    acum = {}
    for mes_data in year.values():
        fyr_mes = (mes_data or {}).get("fyr") or {}
        for clave, vals in fyr_mes.items():
            if not isinstance(vals, dict):
                continue
            a = acum.setdefault(clave, {"fact": 0.0, "rec": 0.0,
                                        "base": 0.0, "kwh": 0.0, "n": 0})
            a["fact"] += vals.get("fact") or 0
            a["rec"] += vals.get("rec") or 0
            a["base"] += vals.get("base") or 0
            a["kwh"] += vals.get("kwh") or 0
            a["n"] += vals.get("n") or 0
    return acum


# Crea una carpeta (año) vacía en el análisis general si no existe
def crear_carpeta_anio(comercializadora, anio):
    if not comercializadora or not anio:
        return False
    folders = dict(comercializadora.analisis_general or {})
    if str(anio) in folders:
        return True
    folders[str(anio)] = {}
    comercializadora.analisis_general = folders
    db.session.commit()
    return True


# Borra la carpeta (año) del análisis general
def eliminar_carpeta_anio(comercializadora, anio):
    if not comercializadora or not anio:
        return False
    folders = dict(comercializadora.analisis_general or {})
    if str(anio) not in folders:
        return False
    del folders[str(anio)]
    comercializadora.analisis_general = folders
    db.session.commit()
    return True


# Devuelve la lista de años que tienen tarifas guardadas (del más nuevo al más viejo)
def tarifas_anios(comercializadora):
    d = comercializadora.tarifas or {}
    return sorted((int(k) for k in d if str(k).isdigit()), reverse=True) if isinstance(d, dict) else []


# Devuelve las tarifas de un año (o del más reciente) y los años disponibles
def get_tarifas(comercializadora, anio=None):
    if not comercializadora:
        return None, []
    d = comercializadora.tarifas
    if not isinstance(d, dict):
        return None, []
    if anio is not None and str(anio) in d:
        return d[str(anio)], tarifas_anios(comercializadora)
    keys = sorted((int(k) for k in d if str(k).isdigit()), reverse=True)
    if keys:
        return d[str(keys[0])], [int(k) for k in keys]
    return None, []


# Guarda las tarifas de un año en la comercializadora
def guardar_tarifas(comercializadora, anio, tarifas_dict):
    if not comercializadora or not anio or not isinstance(tarifas_dict, dict):
        return False
    d = dict(comercializadora.tarifas or {})
    d[str(anio)] = tarifas_dict
    comercializadora.tarifas = d
    db.session.commit()
    return True


# Crea un año de tarifas; si ya existe no hace nada.
# Si hay otro año, copia el último para que sirva de base
def crear_anio_tarifas(comercializadora, anio):
    if not comercializadora or not anio:
        return False
    d = dict(comercializadora.tarifas or {})
    if str(anio) in d:
        return True
    base = None
    keys = sorted((int(k) for k in d if str(k).isdigit()), reverse=True)
    if keys:
        base = d[str(keys[0])]
    d[str(anio)] = json.loads(json.dumps(base)) if base else {
        "residenciales": [], "no_residenciales": [], "autogeneradores": [],
        "actividades": [], "techo": {},
    }
    comercializadora.tarifas = d
    db.session.commit()
    return True


# Borra las tarifas de un año
def eliminar_anio_tarifas(comercializadora, anio):
    if not comercializadora or not anio:
        return False
    d = dict(comercializadora.tarifas or {})
    if str(anio) not in d:
        return False
    del d[str(anio)]
    comercializadora.tarifas = d
    db.session.commit()
    return True


# Devuelve la descripción por grupos si existe (si no, un diccionario vacío)
def get_descripcion(comercializadora):
    if not comercializadora:
        return {}
    d = comercializadora.descripcion
    if isinstance(d, dict) and isinstance(d.get("grupos"), dict):
        return d
    return {}


# Devuelve la cartera acumulada de un año (o la más reciente).
# Se usa deepcopy porque esa información no se debe modificar por accidente
def get_cartera(comercializadora, anio=None):
    if not comercializadora or not comercializadora.cartera_acumulada:
        return None
    d = comercializadora.cartera_acumulada
    if not isinstance(d, dict):
        return None
    if anio is not None and str(anio) in d:
        return copy.deepcopy(d[str(anio)])
    keys = [k for k in d if str(k).isdigit()]
    if not keys:
        return None
    return copy.deepcopy(d[sorted(keys, key=int, reverse=True)[0]])


# Años que tienen cartera guardada (del más nuevo al más viejo)
def carteras_anios(comercializadora):
    d = comercializadora.cartera_acumulada if comercializadora else None
    if not isinstance(d, dict):
        return []
    return sorted((int(k) for k in d if str(k).isdigit()), reverse=True)


# Guarda la cartera de un año en la comercializadora
def save_cartera(comercializadora, anio, cartera_dict, commit=True):
    if not comercializadora or not anio or not isinstance(cartera_dict, dict):
        return False
    d = copy.deepcopy(comercializadora.cartera_acumulada or {})
    d[str(anio)] = copy.deepcopy(cartera_dict)
    comercializadora.cartera_acumulada = d
    if commit:
        db.session.commit()
    return True


# Lee un archivo de referencia (Clientes Especiales) para armar la cartera
# y la guarda; también aprovecha para rellenar las tarifas y la descripción
def cargar_cartera_desde_referencia(comercializadora, path, anio):
    from reports.info_especiales_reader import parse_info_especiales

    data = parse_info_especiales(path)
    cartera = data.get("cartera")
    if not cartera:
        return None
    cartera = json.loads(json.dumps(cartera, default=str))

    # Se recalculan los saldos del año con lo que haya en "meses"
    for c in cartera.get("clientes", []):
        for k in list((c.get("saldos") or {}).keys()):
            if str(k) == str(anio):
                c["saldos"][str(anio)] = 0.0
        if "meses" not in c or not isinstance(c.get("meses"), dict):
            c["meses"] = {}
        c["saldos"][str(anio)] = sum(
            float(c["meses"].get(str(m)) or 0) for m in range(1, 13))
        ma = dict(c.get("meses_anios") or {})
        ma[str(anio)] = dict(c["meses"])
        c["meses_anios"] = ma
        c["meses_2026"] = []
    # Los totales se recalculan con la misma regla del Excel (SUMIF)
    _recalc_totales(cartera)
    save_cartera(comercializadora, anio, cartera, commit=False)

    # Se aprovecha para guardar las tarifas y la descripción que trae el archivo
    if data.get("tarifas"):
        _src = dict(comercializadora.tarifas or {})
        _src[str(anio)] = json.loads(json.dumps(data["tarifas"], default=str))
        comercializadora.tarifas = _src
    if data.get("descripcion"):
        comercializadora.descripcion = json.loads(json.dumps(data["descripcion"], default=str))
    db.session.commit()

    # Se actualiza la cartera de cada reporte ya guardado de ese año
    from models import Report
    reps = (Report.query
            .filter_by(comercializadora_id=comercializadora.id, anio=anio)
            .order_by(Report.mes.asc())
            .all())
    for rep in reps:
        gd = rep.general_data or {}
        esp = rep.especiales_data or {}
        actualizar_cartera_mes(comercializadora, anio, rep.mes,
                               general_data={"facturacion": gd.get("facturacion") or []},
                               especiales_filas=esp.get("filas"))
    return cartera


# Años congelados del archivo de referencia: nunca se escriben por código,
# solo llegan del Excel (2019-2025). El sistema solo llena el año en curso.
ANIO_BASE_CARTERA = 2026


# Crea el año de cartera si no existe (ej. llega enero 2027): clona el año
# más reciente como base, con saldos del año nuevo en cero y meses vacíos.
def ensure_cartera_anio(comercializadora, anio):
    if not comercializadora or not anio or int(anio) < ANIO_BASE_CARTERA:
        return False
    d = dict(comercializadora.cartera_acumulada or {})
    if str(anio) in d:
        return True
    keys = [k for k in d if str(k).isdigit()]
    if not keys:
        return False
    base = copy.deepcopy(d[sorted(keys, key=int, reverse=True)[0]])
    for c in base.get("clientes", []):
        saldos = dict(c.get("saldos") or {})
        saldos[str(anio)] = 0.0
        c["saldos"] = saldos
        ma = dict(c.get("meses_anios") or {})
        if "2026" not in ma and isinstance(c.get("meses"), dict):
            ma["2026"] = dict(c["meses"])
        ma[str(anio)] = {}
        c["meses_anios"] = ma
    for clave in ("total_favor", "total_pagar"):
        t = dict(base.get(clave) or {})
        t[str(anio)] = 0.0
        base[clave] = t
    d[str(anio)] = base
    comercializadora.cartera_acumulada = d
    db.session.commit()
    return True


# Total de un cliente (columna L del Excel: suma de sus años sin el 2019
# prescrito) para el VALOR TOTAL CARTERA
def _total_cliente(c):
    saldos = c.get("saldos") or {}
    return sum(float(saldos.get(k) or 0) for k in saldos
               if str(k).isdigit() and str(k) != "2019")


# Recalcula favor/pagar por año (SUMIF por columna, como filas 60-61) y los
# totales generales como SUMIF sobre el total de cada cliente (columna L)
def _recalc_totales(cartera):
    clientes = cartera.get("clientes", [])
    years = sorted({str(k) for c in clientes for k in (c.get("saldos") or {})
                    if str(k).isdigit()}, key=int)
    d_f = dict(cartera.get("total_favor") or {})
    d_p = dict(cartera.get("total_pagar") or {})
    for y in years:
        fav = pag = 0.0
        for c in clientes:
            v = float((c.get("saldos") or {}).get(y) or 0)
            if v > 0:
                fav += v
            elif v < 0:
                pag += v
        d_f[y] = fav
        d_p[y] = pag
    tot_f = tot_p = 0.0
    for c in clientes:
        t = _total_cliente(c)
        if t > 0:
            tot_f += t
        elif t < 0:
            tot_p += t
    d_f["total"] = tot_f
    d_p["total"] = tot_p
    cartera["total_favor"] = d_f
    cartera["total_pagar"] = d_p
    return cartera


# Después de procesar un mes, guarda el saldo de ese mes en la cartera.
# Igual que el VLOOKUP del Excel de referencia: el mes trae el R-L
# (recaudo - liquidación) del informe especial mensual ya hecho,
# buscado por ID de cuenta.
# Solo años en curso (2026 en adelante): 2019-2025 quedan congelados.
# La columna del año es la suma de sus meses, como en el Excel.
# El grupo/tarifa/periodo se mantienen desde la cartera (edición manual,
# como el "CAMBIAR DE GRUPO" del Excel); el nombre sí se sincroniza porque
# los archivos mandan.
def actualizar_cartera_mes(comercializadora, anio, mes, general_data=None,
                           especiales_filas=None):
    if not anio or int(anio) < ANIO_BASE_CARTERA or not mes:
        return False
    anio = int(anio)
    mes = int(mes)
    if not (1 <= mes <= 12):
        return False
    ensure_cartera_anio(comercializadora, anio)
    cartera = get_cartera(comercializadora, anio)
    if not cartera:
        return False

    # Migra los meses viejos al formato por año (los "meses" planos son 2026)
    for c in cartera.get("clientes", []):
        ma = c.get("meses_anios")
        if not isinstance(ma, dict):
            ma = {}
        if "2026" not in ma and isinstance(c.get("meses"), dict):
            ma["2026"] = dict(c["meses"])
        ma.setdefault(str(anio), {})
        c["meses_anios"] = ma

    # Mapa id -> {"saldo": R-L del mes, "nombre": nombre del archivo}
    saldos_mes = {}
    if especiales_filas:
        for e in especiales_filas:
            clave = str(e.get("id") or "").strip()
            if not clave:
                continue
            saldos_mes[clave] = {
                "saldo": float(e.get("val_rec") or 0),
                "nombre": str(e.get("nombre") or "").strip(),
            }
    elif general_data:
        # Respaldo para informes viejos sin especiales guardados: aproxima el
        # R-L con lo facturado/recaudado del mes y la tarifa de cartera
        facturacion = general_data.get("facturacion") or []
        recaudo = general_data.get("recaudo") or []
        if not facturacion and not recaudo:
            return False
        por_id = {}
        for r in facturacion:
            clave = str(r.get("Id_cliente") or "").strip()
            if not clave:
                continue
            d = por_id.setdefault(clave, {"fact": 0.0, "cons": 0.0, "rec": 0.0})
            d["fact"] += float(r.get("ImpuestoAP") or 0)
            d["cons"] += float(r.get("Consumo$") or 0)
        for r in recaudo:
            clave = str(r.get("Id_cliente") or "").strip()
            if not clave:
                continue
            d = por_id.setdefault(clave, {"fact": 0.0, "cons": 0.0, "rec": 0.0})
            d["rec"] += float(r.get("ValorAP") or 0)
        tarifas = {str(c.get("id")).strip(): float(c.get("tarifa_min") or 0)
                   for c in cartera.get("clientes", [])}
        TEC = 35018100
        for clave, d in por_id.items():
            IAP = min(max(tarifas.get(clave, 0.0), d["cons"] * 0.10), TEC)
            saldos_mes[clave] = {"saldo": d["rec"] - IAP, "nombre": ""}
    else:
        return False

    if not saldos_mes:
        return False

    claves_idx = {str(c["id"]).strip(): i for i, c in enumerate(cartera.get("clientes", []))}
    guardado = False
    # Se escribe el mes (siempre se sobrescribe: como el VLOOKUP, la cartera
    # refleja lo último del informe mensual, así se puede reprocesar un mes)
    for clave, i in claves_idx.items():
        if clave not in saldos_mes:
            continue
        cl = cartera["clientes"][i]
        if saldos_mes[clave]["nombre"]:
            cl["nombre"] = saldos_mes[clave]["nombre"]
        ma = cl["meses_anios"]
        md = dict(ma.get(str(anio)) or {})
        md[str(mes)] = float(saldos_mes[clave]["saldo"])
        ma[str(anio)] = md
        if anio == 2026:
            cl["meses"] = dict(md)
            meses = set(cl.get("meses_2026") or [])
            meses.add(mes)
            cl["meses_2026"] = sorted(int(x) for x in meses)
        saldos = cl.get("saldos") or {}
        saldos[str(anio)] = sum(float(md.get(str(m)) or 0) for m in range(1, 13))
        cl["saldos"] = saldos
        cl["total_cartera"] = _total_cliente(cl)
        guardado = True

    if guardado:
        _recalc_totales(cartera)
        save_cartera(comercializadora, anio, cartera)
    return guardado


# Guarda o actualiza un reporte en la base de datos.
# Si ya existe un reporte del mismo municipio/comercializadora/año/mes, lo actualiza
def save_report_to_db(municipio, comercializadora, anio, mes,
                       general_data, especiales_data,
                       fact_path=None, rec_path=None,
                       analisis_ejecutivo=None, novedades=None):
    m = get_or_create_municipio(municipio)
    c = get_or_create_comerc(m.id, comercializadora)

    # Normaliza: textos vacíos ('null', '{}', ...) equivalen a no tener datos
    if isinstance(especiales_data, str) and especiales_data.strip() in ('null', '{}', '[]', ''):
        especiales_data = None
    if isinstance(general_data, str) and general_data.strip() in ('null', '{}', '[]', ''):
        general_data = None

    existing = Report.query.filter_by(comercializadora_id=c.id, anio=anio, mes=mes).first()
    if existing:
        existing.general_data = general_data
        existing.especiales_data = especiales_data
        existing.archivo_fact_path = fact_path
        existing.archivo_rec_path = rec_path
        existing.status = "procesado"
        existing.updated_at = ahora_utc()
        if analisis_ejecutivo is not None:
            existing.analisis_ejecutivo = analisis_ejecutivo
        if novedades is not None:
            existing.novedades = novedades
        report = existing
    else:
        report = Report(
            comercializadora_id=c.id, anio=anio, mes=mes,
            general_data=general_data, especiales_data=especiales_data,
            archivo_fact_path=fact_path, archivo_rec_path=rec_path,
            status="procesado",
            analisis_ejecutivo=analisis_ejecutivo,
            novedades=novedades,
        )
        db.session.add(report)

    db.session.commit()

    # Se precalculan las alertas/anomalías del informe (idea #2) y el
    # resumen numérico mensual (ideas #3, #6, #9) para que la página y el
    # chat respondan al instante. Si falla, no interrumpe la carga.
    try:
        from core.alertas import guardar_alertas
        guardar_alertas(report)
    except Exception:
        db.session.rollback()
    try:
        from core.ia_chat import guardar_resumen
        guardar_resumen(report)
    except Exception:
        db.session.rollback()

    return report


# Arma el texto del análisis ejecutivo para el informe.
# Explica en palabras cuánto se dejó de facturar y de percibir
def texto_ejecutivo(municipio, comercializadora, anio, mes, data, variacion=None):
    a = data.get("analisis", {})
    periodo = f"{MESES[mes] if 1 <= mes <= 12 else mes} {anio}"

    iap_df = a.get("iap_dejado_facturar", 0) or 0
    iap_mun = a.get("iap_no_percibido_municipio", 0) or 0
    millones_agente = round((iap_df or 0) / 1_000_000, 1)
    millones_mun = round((iap_mun or 0) / 1_000_000, 1)

    # Se elige la frase según la variación que hubo con el periodo anterior.
    # La variación es (anterior - actual) / anterior: positiva = el saldo
    # negativo bajó (mejora), negativa = el saldo negativo subió (desmejora)
    if variacion is None:
        frase_var = "en comparaci\u00f3n con el periodo anterior"
    elif variacion > 0:
        frase_var = f"disminuy\u00f3 {variacion:,.1f}% en comparaci\u00f3n con el anterior, lo que indica una mejora en la aplicaci\u00f3n de tarifas del IAP"
    elif variacion < 0:
        frase_var = f"se increment\u00f3 {abs(variacion):,.1f}% en comparaci\u00f3n con el anterior, lo que indica que el agente de recaudo desmejor\u00f3 la aplicaci\u00f3n de tarifas del IAP"
    else:
        frase_var = "se mantuvo sin variaci\u00f3n en comparaci\u00f3n con el anterior"

    # Se cuentan y nombran (hasta 5) las cuentas que no facturaron
    filas_txt = ""
    sin_f = a.get("list_sin_facturar", [])
    if sin_f:
        nombres = ", ".join(x["nombre"] for x in sin_f[:5])
        if len(sin_f) > 5:
            nombres += f" y {len(sin_f) - 5} m\u00e1s"
        filas_txt += (f" Se identificaron {a.get('cuentas_sin_facturar', 0)} cuentas sin "
                      f"facturaci\u00f3n ({sin_f[0]['estado_fact']}): {nombres}.")

    text = (
        f"De acuerdo con el an\u00e1lisis de {periodo}, se estima que el agente de recaudo "
        f"est\u00e1 dejando de facturar hasta ${millones_agente:,.1f} Millones, y el Municipio "
        f"de {municipio} dej\u00f3 de percibir hasta ${millones_mun:,.1f} millones por concepto "
        f"de impuesto de alumbrado p\u00fablico correspondiente a la facturaci\u00f3n de cuentas "
        f"especiales. Para este periodo, el saldo negativo {frase_var}.{filas_txt}"
    )
    return text


# Calcula el % de variación del saldo negativo de un reporte
# comparándolo con el reporte anterior de la misma comercializadora.
# Fórmula: (saldo anterior - saldo actual) / saldo anterior * 100
# (con los saldos en magnitud positiva). Positiva = el saldo negativo bajó
# (mejora); negativa = el saldo negativo subió (desmejora).
def variacion_saldo_neg(report):
    if not report or not report.especiales_data:
        return None
    actual = ((report.especiales_data.get("totales") or {}).get("saldo_neg") or 0) * -1
    if actual == 0:
        return None

    # Se buscan los reportes de la misma comercializadora y se toma el anterior
    prev = (Report.query
            .filter(Report.comercializadora_id == report.comercializadora_id,
                    Report.id != report.id)
            .order_by(Report.anio.desc(), Report.mes.desc())
            .all())
    prev = [p for p in prev if (p.anio, p.mes) < (report.anio, report.mes) and p.especiales_data]
    if not prev:
        return None
    anterior = ((prev[0].especiales_data.get("totales") or {}).get("saldo_neg") or 0) * -1
    if anterior == 0:
        return None
    return (anterior - actual) / anterior * 100
