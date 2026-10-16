# -*- coding: utf-8 -*-
"""Informe exógeno de EMSA (Villavicencio): un solo archivo con todo.

La hoja "Exogena" trae por usuario lo facturado, lo recaudado y la cartera.
Desde ella se arman los dos archivos lógicos (facturación y recaudo) que
entiende el sistema, para Villavicencio se sube SOLO este archivo.
"""

import pandas as pd

MUNICIPIO = "Villavicencio"
COMERCIALIZADORA = "EMSA"
EMPRESA_NOMBRE = "ELECTRIFICADORA DEL META S.A. (EMSA)"

# Nombres de columnas que se buscan por nombre (no por posición, porque
# el formato del exógeno cambia: a veces trae BARRIO/CORREO, a veces no)
_NOMBRES_COL = {
    "id": "NUMERO DE IDENTIFICACION USUARIO",
    "mercado": "MERCADO (REGULADO O NO REGULADO)",
    "sector": "SECTOR (URBANO – RURAL)",
    "destino": "DESTINO ECONOMICO DEL PREDIO",
    "estrato": "ESTRATO SOCIOECONOMICO",
    "factura": "NUMERO DE FACTURA",
    "periodo": "PERIODO / MES",
    "kwh": "CONSUMO DE KILOVATIOS (KW/H)",
    "cons_valor": "VALOR CONSUMO DE ENERGIA (BASE PARA FACTURACION)",
    "imp_fact": "VALOR IMPUESTO DE ALUMBRADO PUBLICO FACTURADO",
    "imp_rec": "VALOR IMPUESTO DE ALUMBRADO PUBLICO RECAUDADO",
    "apellido1": "APELLIDO 1",
    "apellido2": "APELLIDO 2",
    "nombre1": "NOMBRE 1",
    "nombre2": "NOMBRE 2",
    "razon": "RAZON SOCIAL",
    "direccion": "DIRECCION DE SUMINISTRO DEL SERVICIO",
}


def _resolver_columnas(raw):
    """Devuelve un dict con el índice (0-based) de cada columna por nombre.
    Si un nombre no se encuentra, su valor es None."""
    cols = {str(c).strip().upper(): i for i, c in enumerate(raw.columns)}
    return {k: cols.get(v.strip().upper()) for k, v in _NOMBRES_COL.items()}

# Tarifas del Acuerdo 515 de 12/04/2021 (Art. 242): % sobre la LCEE
TARIFAS_LCEE_2026 = {
    "esquema": "lcee",
    "lcee": {
        "fuente": ("Acuerdo No. 515 de 12/04/2021 - Concejo Municipal de "
                   "Villavicencio (Capitulo XVI, Impuesto de Alumbrado Publico)"),
        "base_gravable": ("Liquidacion del Consumo de Energia Electrica (LCEE) "
                          "antes de contribuciones o subsidios (prepago o pospago); "
                          "el avaluo catastral en los casos pertinentes."),
        "residenciales": [
            {"estrato": 1, "porcentaje": 1},
            {"estrato": 2, "porcentaje": 2},
            {"estrato": 3, "porcentaje": 3},
            {"estrato": 4, "porcentaje": 8},
            {"estrato": 5, "porcentaje": 12},
            {"estrato": 6, "porcentaje": 12},
        ],
        "sectores": [
            {"sector": "Comercial", "porcentaje": 15},
            {"sector": "Industrial", "porcentaje": 15},
            {"sector": "Oficial", "porcentaje": 15},
            {"sector": "Temporales", "porcentaje": 15},
            {"sector": "Autogeneradores", "porcentaje": 15},
        ],
        "predios_sin_servicio": {
            "por_mil": 1,
            "base": "avaluo que sirve de base para liquidar el impuesto predial",
        },
        "notas": [
            "Art. 235: autorizacion legal (Ley 97 de 1913, Ley 84 de 1915 y Ley 1819 de 2016).",
            "Art. 240: base gravable = LCEE antes de contribucion o subsidios; incluye "
            "energia alternativa, cogenerada y autogenerada.",
            "Art. 241: el impuesto se causa mensualmente.",
            "Art. 242 (Parag. 2): contribuyentes no clasificados (predios y/o lotes sin "
            "servicio) pagan 1 por mil sobre el avaluo.",
            "Art. 243: los autogeneradores pagan segun la LCEE; si ademas son usuarios, "
            "se suma el consumo facturado mas la energia autogenerada.",
            "Art. 246: responsables del recaudo son las empresas prestadoras, "
            "distribuidores y comercializadores de energia; el impuesto se recauda "
            "conjuntamente con el servicio e incorporado en la factura.",
            "Art. 249: destinacion exclusiva a la prestacion, mejora, modernizacion y "
            "ampliacion del alumbrado publico.",
        ],
    },
}


def _str(v):
    """Convierte una celda a texto limpio ('' si está vacía)."""
    if v is None:
        return ""
    if isinstance(v, float) and (pd.isna(v) or v != v):
        return ""
    s = str(v).strip()
    if s.lower() in ("nan", "none", "nat"):
        return ""
    return s


def _num(v):
    """Convierte una celda a número (None si no es numérica)."""
    s = _str(v).replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _id(v):
    """Identificación del usuario en número entero cuando sea posible."""
    n = _num(v)
    if n is None:
        return None
    return int(n) if float(n).is_integer() else n


def _nombre(nombre1, nombre2, apellido1, apellido2, razon):
    """Arma el nombre del contribuyente: razón social o nombres + apellidos."""
    if razon:
        return razon
    partes = [p for p in (apellido1, apellido2, nombre1, nombre2) if p]
    return " ".join(partes)


def es_archivo_emsa(path):
    """Dice si el Excel es el exógeno de EMSA (tiene hoja Exogena)."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return any((s or "").strip().lower() == "exogena" for s in wb.sheetnames)
        finally:
            wb.close()
    except Exception:
        return False


def leer_exogena_emisa(path):
    """Lee el archivo exógeno de EMSA y devuelve los DataFrames de
    facturación y recaudo en el formato interno que espera el sistema."""
    # Las 3 primeras filas son metadatos (Nit, nombre, periodo); la fila 4 es el encabezado
    from core.utils import leer_excel
    raw = leer_excel(path, sheet_name="Exogena", header=3, dtype=str)

    # Se resuelven las columnas por nombre (el formato cambia entre archivos)
    C = _resolver_columnas(raw)

    def _col(clave):
        """Devuelve la serie de la columna, o Series vacía si no existe."""
        idx = C.get(clave)
        if idx is None:
            return pd.Series([None] * len(raw), index=raw.index)
        return raw.iloc[:, idx]

    ids = _col("id").map(_id)
    mercado = _col("mercado").map(_str)
    sector = _col("sector").map(_str)
    destino = _col("destino").map(_str)
    estrato = _col("estrato").map(_str)
    direccion = _col("direccion").map(_str)

    nombres = []
    for t in raw.itertuples(index=False, name=None):
        nombres.append(_nombre(
            t[C["nombre1"]] if C.get("nombre1") is not None else None,
            t[C["nombre2"]] if C.get("nombre2") is not None else None,
            t[C["apellido1"]] if C.get("apellido1") is not None else None,
            t[C["apellido2"]] if C.get("apellido2") is not None else None,
            t[C["razon"]] if C.get("razon") is not None else None,
        ))

    fact = pd.DataFrame({
        "Id_cliente": ids,
        "numero_factura": _col("factura").map(_str),
        "monto_AP": _col("imp_fact").map(lambda v: _num(v) or 0.0),
        "monto_kwh": _col("cons_valor").map(_num),
        "consumo_kwh": _col("kwh").map(_num),
        "Empresa": EMPRESA_NOMBRE,
        "Nombre": nombres,
        "dirección": direccion,
        "desc_clase": destino,
        "mercado": mercado,
        "estrato_socioeconomico": estrato,
        "Sector": sector,
        "Periodicidad": "Mensual",
    })

    rec = pd.DataFrame({
        "Id_cliente": ids,
        "valor_recaudado": _col("imp_rec").map(lambda v: _num(v) or 0.0),
        "Empresa": EMPRESA_NOMBRE,
        "mercado": mercado,
        "nombre": nombres,
        "direccion": direccion,
        "desc_clase": destino,
        "Sector": sector,
        "estrato_socioeconomico": estrato,
    })

    # Se descartan las filas sin identificación de usuario
    fact = fact[fact["Id_cliente"].notna()].copy()
    rec = rec[rec["Id_cliente"].notna()].copy()

    # Se intenta sacar el periodo desde la columna "PERIODO / MES" (ej. "2026/6")
    periodos = [_str(v) for v in _col("periodo").tolist()]
    periodos = [p for p in periodos if "/" in p]
    periodo = None
    if periodos:
        try:
            from collections import Counter
            cand = Counter(periodos).most_common(1)[0][0]
            anio, mes = cand.split("/")
            periodo = (int(anio), int(mes))
        except (ValueError, TypeError):
            periodo = None

    return fact, rec, periodo
